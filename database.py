import pickle
import io
from pathlib import Path
import numpy as np
from sklearn.preprocessing import normalize


class PartDatabase:
    def __init__(self, excel_path=None):
        self.excel_path = Path(excel_path) if excel_path else None
        self.records = []  # list of {料號, 名稱, 庫存, features}
        self.features = None

    def build(self, extractor):
        if self.excel_path is None:
            raise ValueError("請指定 Excel 檔案路徑")

        import openpyxl
        wb = openpyxl.load_workbook(self.excel_path)
        ws = wb.active

        parts_left = {}
        parts_right = {}
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, 1).value
            if v and str(v).strip():
                parts_left[r] = str(v).strip()
            v = ws.cell(r, 9).value
            if v and str(v).strip():
                parts_right[r] = str(v).strip()

        images = ws._images
        if not images:
            print("! Excel 中找不到任何嵌入圖片")
            return

        print(f"從 Excel 讀取到 {len(images)} 張圖片")

        records_map = {}  # 料號 -> {料號, 名稱, 庫存, features}

        for img in images:
            r = img.anchor._from.row + 1
            c = img.anchor._from.col + 1

            part_no = ''
            side = ''
            if c <= 8:
                for pr in sorted(parts_left.keys(), reverse=True):
                    if pr < r:
                        part_no = parts_left[pr]
                        side = 'left'
                        break
            else:
                for pr in sorted(parts_right.keys(), reverse=True):
                    if pr < r:
                        part_no = parts_right[pr]
                        side = 'right'
                        break

            if not part_no:
                continue

            raw = img._data()
            if callable(raw):
                raw = raw()
            pil_img = _bytes_to_pil(raw)
            if pil_img is None:
                continue

            try:
                feat = extractor.extract(pil_img)
            except Exception as e:
                print(f"! 處理 {part_no} 圖片失敗: {e}")
                continue

            if part_no not in records_map:
                name_col = 2 if side == 'left' else 10
                name = ws.cell(r - 1, name_col).value or ''
                records_map[part_no] = {
                    '料號': part_no,
                    '名稱': str(name).strip(),
                    '庫存': '',
                    'features': [],
                }

            records_map[part_no]['features'].append(feat)

        self.records = list(records_map.values())
        if not self.records:
            print("! 沒有成功建立任何資料")
            self.features = np.empty((0, 512), dtype=np.float32)
            return

        all_feats = []
        for rec in self.records:
            avg_feat = np.mean(rec['features'], axis=0).astype(np.float32)
            all_feats.append(avg_feat)
            del rec['features']

        self.features = normalize(np.array(all_feats, dtype=np.float32))
        print(f"[OK] 成功建立 {len(self.records)} 筆零件資料")

    def save(self, path):
        data = {
            'records': self.records,
            'features': self.features,
        }
        with open(path, 'wb') as f:
            pickle.dump(data, f)

    def load(self, path):
        with open(path, 'rb') as f:
            data = pickle.load(f)
        self.records = data['records']
        self.features = data['features']

    def search(self, query_features, top_k=5):
        if self.features is None or len(self.records) == 0:
            return []
        query_norm = normalize(query_features.reshape(1, -1))
        sims = np.dot(self.features, query_norm.T).flatten()
        top_idx = np.argsort(sims)[::-1][:top_k]
        results = []
        for idx in top_idx:
            if sims[idx] < 0.1:
                continue
            rec = self.records[idx]
            results.append({
                'similarity': round(float(sims[idx]) * 100, 1),
                '料號': rec.get('料號', ''),
                '名稱': rec.get('名稱', ''),
                '庫存': rec.get('庫存', ''),
            })
        return results


def _bytes_to_pil(data):
    try:
        from PIL import Image
        if isinstance(data, bytes):
            return Image.open(io.BytesIO(data))
        return Image.open(io.BytesIO(data.read()))
    except Exception:
        return None
